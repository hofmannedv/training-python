# Demonstrate usage of OpenPyxl to identify sheetnames from an 
# Excel spreadsheet
#
# (C) 2024 Frank Hofmann, Freiburg, Germany <info@efho.de>
# Published under GNU Public License (GPL)
#
# Inspired by examples from:
# https://realpython.com/openpyxl-excel-spreadsheets-python/

from openpyxl import Workbook, load_workbook

# define Excel file to read, and load it
filename = "hello_world.xlsx"
workbook = load_workbook(filename)

print("This file contains the following sheets:")
for sheet in workbook.sheetnames:
    print(sheet)

print(" ")
print("Adding spreadsheet Fruits to it ...")

# add sheet Fruits
workbook.create_sheet("Fruits")

print(" ")
print("This file now contains the following sheets:")
for sheet in workbook.sheetnames:
    print(sheet)

# save file as hello_world.xlsx
workbook.save(filename="hello_world.xlsx")
